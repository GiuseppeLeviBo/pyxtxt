import io

import pytest

from pyxtxt import xtxt
from pyxtxt.estrattori import estrattori

RTF_SOURCE = r"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Hello {\b RTF} world.\par}"
TEX_SOURCE = "\\documentclass{article}\n\\begin{document}\n\\section{Intro}\nHello \\textbf{LaTeX}.\n\\end{document}\n"
MD_SOURCE = "# Title\n\nSome **bold** text\n\n- item\n"


def test_plain_text_from_bytes():
    assert xtxt(b"just some plain text") == "just some plain text"


def test_missing_file_returns_none(tmp_path):
    assert xtxt(str(tmp_path / "does-not-exist.txt")) is None


def test_unsupported_type_returns_none():
    assert xtxt(b"\x00\x01\x02\x03" * 64) is None


def test_markdown_file(tmp_path):
    pytest.importorskip("markdown")
    pytest.importorskip("bs4")
    path = tmp_path / "doc.md"
    path.write_text(MD_SOURCE)
    text = xtxt(str(path))
    assert "Title" in text and "bold" in text
    assert "**" not in text and "#" not in text


def test_markdown_buffer_uses_name_hint():
    pytest.importorskip("markdown")
    pytest.importorskip("bs4")
    buffer = io.BytesIO(MD_SOURCE.encode())
    buffer.name = "notes.md"
    text = xtxt(buffer)
    assert "**" not in text and "Title" in text


def test_rtf_file(tmp_path):
    pytest.importorskip("striprtf")
    path = tmp_path / "doc.rtf"
    path.write_text(RTF_SOURCE)
    text = xtxt(str(path))
    assert "\\rtf1" not in text
    assert "Hello RTF world." in text


def test_latex_file(tmp_path):
    pytest.importorskip("pylatexenc")
    path = tmp_path / "doc.tex"
    path.write_text(TEX_SOURCE)
    text = xtxt(str(path))
    assert "\\documentclass" not in text
    assert "LaTeX" in text


def test_corrupted_pdf_returns_none_not_string(tmp_path):
    pytest.importorskip("pymupdf")
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4\n garbage")
    assert xtxt(str(path)) is None


def test_pdf_text(tmp_path):
    fitz = pytest.importorskip("pymupdf")
    path = tmp_path / "doc.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Hello PDF")
    doc.save(str(path))
    assert "Hello PDF" in xtxt(str(path))


def test_xlsx_is_not_truncated(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    path = tmp_path / "big.xlsx"
    wb = openpyxl.Workbook()
    for i in range(500):
        wb.active.append([i, f"row{i}"])
    wb.save(str(path))
    text = xtxt(str(path))
    assert "0 | row0" in text
    assert "499 | row499" in text


def test_xlsx_row_limit_is_still_available(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from pyxtxt.estrattori.xlsx import xtxt_xlsx

    wb = openpyxl.Workbook()
    for i in range(50):
        wb.active.append([i])
    buffer = io.BytesIO()
    wb.save(buffer)
    text = xtxt_xlsx(buffer, max_rows_per_sheet=10)
    assert "9" in text.splitlines() and "10" not in text.splitlines()


def test_docx_text(tmp_path):
    docx = pytest.importorskip("docx")
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_paragraph("Hello DOCX")
    document.save(str(path))
    assert "Hello DOCX" in xtxt(str(path))


def test_html_text():
    pytest.importorskip("bs4")
    text = xtxt(b"<html><body><p>Visible paragraph</p></body></html>")
    assert "Visible paragraph" in text and "<p>" not in text


def test_ollama_ocr_takes_precedence_over_easyocr():
    """Extractor modules load in sorted order, so ocr_ollama always overrides ocr."""
    pytest.importorskip("ollama")
    pytest.importorskip("PIL")
    assert estrattori["image/png"].__module__.endswith("ocr_ollama")


# --- Inputs -----------------------------------------------------------------

def test_binary_file_object(tmp_path):
    pytest.importorskip("markdown")
    path = tmp_path / "doc.md"
    path.write_text(MD_SOURCE)
    with open(path, "rb") as f:
        text = xtxt(f)
    assert "Title" in text and "**" not in text


def test_text_mode_file_object_returns_none(tmp_path):
    path = tmp_path / "doc.txt"
    path.write_text("hello")
    with open(path) as f:
        assert xtxt(f) is None


def test_xlsx_from_bytes_is_detected(tmp_path):
    """libmagic needs more than the first 2 KB to recognise XLSX."""
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    wb.active.append(["cell from bytes"])
    buffer = io.BytesIO()
    wb.save(buffer)
    assert "cell from bytes" in xtxt(buffer.getvalue())


# --- Document formats -------------------------------------------------------

def test_docx_tables_headers_and_footers():
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Header text"
    document.sections[0].footer.paragraphs[0].text = "Footer text"
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "Merged"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    table.cell(1, 2).add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
    document.add_paragraph("After table")
    buffer = io.BytesIO()
    document.save(buffer)

    lines = xtxt(buffer.getvalue()).splitlines()
    assert lines == ["Header text", "Before table", "A1 | Merged", "A2 | B2 | Nested", "After table", "Footer text"]


def test_pptx_groups_tables_and_notes():
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Slide title"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text_frame.text = "Inside group"
    table = slide.shapes.add_table(1, 2, 0, 0, Inches(4), Inches(1)).table
    table.cell(0, 0).text = "H1"
    table.cell(0, 1).text = "H2"
    slide.notes_slide.notes_text_frame.text = "Speaker notes"
    buffer = io.BytesIO()
    presentation.save(buffer)

    assert xtxt(buffer.getvalue()).splitlines() == ["Slide title", "Inside group", "H1 | H2", "Speaker notes"]


def test_odt_headings_spans_and_lists():
    pytest.importorskip("odf")
    from odf.opendocument import OpenDocumentText
    from odf.text import H, List, ListItem, P, Span

    document = OpenDocumentText()
    document.text.addElement(H(outlinelevel=1, text="Heading"))
    paragraph = P(text="Plain ")
    paragraph.addElement(Span(text="spanned"))
    paragraph.addText(" tail")
    document.text.addElement(paragraph)
    items = List()
    item = ListItem()
    item.addElement(P(text="List item"))
    items.addElement(item)
    document.text.addElement(items)
    buffer = io.BytesIO()
    document.save(buffer)

    assert xtxt(buffer.getvalue()).splitlines() == ["Heading", "Plain spanned tail", "List item"]


def test_svg_tspan_and_empty_text():
    pytest.importorskip("lxml")
    svg = (
        b'<svg xmlns="http://www.w3.org/2000/svg">'
        b"<text>Hi</text><text/><text><tspan>nested</tspan> <tspan>words</tspan></text>"
        b"</svg>"
    )
    assert xtxt(svg).splitlines() == ["Hi", "nested words"]


# --- EXIF and email ---------------------------------------------------------

@pytest.mark.parametrize("image_format", ["JPEG", "PNG", "WEBP"])
def test_exif_formats_rationals_and_gps(image_format):
    pytest.importorskip("PIL")
    from PIL import Image
    from PIL.TiffImagePlugin import IFDRational
    from pyxtxt import xtxt_exif

    exif = Image.Exif()
    exif[0x010F] = "Canon"
    camera = exif.get_ifd(0x8769)
    camera[0x829D] = IFDRational(28, 10)   # FNumber
    camera[0x829A] = IFDRational(1, 250)   # ExposureTime
    camera[0x920A] = IFDRational(50, 1)    # FocalLength
    gps = exif.get_ifd(0x8825)
    gps[1] = "N"
    gps[2] = (IFDRational(44, 1), IFDRational(30, 1), IFDRational(0, 1))
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8)).save(buffer, image_format, exif=exif)
    buffer.seek(0)

    text = xtxt_exif(buffer)
    for expected in ["Make: Canon", "Aperture: f/2.8", "Shutter Speed: 1/250s", "Focal Length: 50mm", "GPS Latitude: 44.500000° N"]:
        assert expected in text
    assert "ExifOffset" not in text


def test_no_fake_exif_mime_types():
    from pyxtxt import extxt_available_formats

    assert not [mime for mime in extxt_available_formats() if "+exif" in mime]


class _FakeMsg:
    def __init__(self, body, html_body):
        self.body = body
        self.htmlBody = html_body
        self.closed = False

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.closed = True


@pytest.mark.parametrize(
    "body, html_body, expected",
    [
        ("Plain body\r\n", b"<p>Plain body</p>", "Plain body"),
        ("", b"<html><body><p>Only HTML</p></body></html>", "Only HTML"),
    ],
    ids=["plain-text-preferred", "html-fallback"],
)
def test_msg_body_and_close(monkeypatch, body, html_body, expected):
    pytest.importorskip("extract_msg")
    pytest.importorskip("bs4")
    from pyxtxt.estrattori import msg as msg_module

    fake = _FakeMsg(body, html_body)
    monkeypatch.setattr(msg_module.extract_msg, "openMsg", lambda content: fake)
    assert msg_module.xtxt_msg(io.BytesIO(b"fake msg bytes")) == expected
    assert fake.closed


# --- EML --------------------------------------------------------------------

def _email_with_alternative_and_attachment():
    from email.message import EmailMessage

    message = EmailMessage()
    message["Return-Path"] = "<sender@example.com>"
    message["From"] = "sender@example.com"
    message["Subject"] = "Report"
    message.set_content("Plain version of the body")
    message.add_alternative("<html><body><p>HTML version of the body</p></body></html>", subtype="html")
    message.add_attachment(b"attached text file", maintype="text", subtype="plain", filename="notes.txt")
    return message.as_bytes()


def test_eml_plain_body_once_without_attachments():
    pytest.importorskip("bs4")
    assert xtxt(_email_with_alternative_and_attachment()) == "Plain version of the body"


def test_eml_html_only_and_extension_hint(tmp_path):
    pytest.importorskip("bs4")
    from email.message import EmailMessage

    message = EmailMessage()
    message["Subject"] = "Starts with Subject, so libmagic sees plain text"
    message.set_content("<p>Only <b>HTML</b> here</p>", subtype="html")
    path = tmp_path / "message.eml"
    path.write_bytes(message.as_bytes())
    text = xtxt(str(path))
    assert "Only" in text and "<p>" not in text and "Subject:" not in text
