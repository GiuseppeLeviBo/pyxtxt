from . import register_extractor
import io
import zipfile
try:
    from pptx import Presentation
    from pptx.shapes.group import GroupShape
except ImportError:
    Presentation = None

if Presentation:
    def _shape_lines(shape):
        """Yield the text of a shape, descending into groups and tables."""
        if isinstance(shape, GroupShape):
            for child in shape.shapes:
                yield from _shape_lines(child)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    yield " | ".join(cells)
        elif getattr(shape, "has_text_frame", False) and shape.text_frame.text:
            yield shape.text_frame.text

    def xtxt_pptx(file_buffer) -> str:
        try:
            # Convertiamo il file_buffer (che è già un BytesIO o simile) in modo da poterlo riusare
            file_buffer.seek(0)
            data = file_buffer.read()
            buffer_copy = io.BytesIO(data)

            if not zipfile.is_zipfile(buffer_copy):
                print("⚠️  Invalid PPTX (not a ZIP file)")
                return ""

            # Se è un file zip valido, possiamo ripassare i dati a Presentation
            buffer_copy.seek(0)
            prs = Presentation(buffer_copy)

            lines = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    lines.extend(_shape_lines(shape))
                # Speaker notes
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame
                    if notes is not None and notes.text:
                        lines.append(notes.text)
            return "\n".join(lines)

        except Exception as e:
            print(f"⚠️ Error during PPTX extraction: {e}")
            return ""

    register_extractor(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        xtxt_pptx,
        name="PPTX"
    )
